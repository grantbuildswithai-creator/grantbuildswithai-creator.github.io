"""Turns a guide's Markdown (guides/<slug>.md) into the Word version people can download.

    python tools/guide_to_docx.py build-your-own-app "Build Your Own App with AI" static/reports/build-your-own-app-guide.docx

Uses Word's built-in styles (Title, Heading 2/3, List Bullet, Quote), like the other guide downloads.
"""
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent.parent
SITE = "https://grantbuildswithai.com"
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_link(paragraph, text, url):
    rid = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        elif part.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            if url.startswith("../"):
                url = SITE + "/" + url.lstrip("./")
            add_link(paragraph, label, url)
        else:
            paragraph.add_run(part)


def convert(slug, title, out):
    md = (HERE / "guides" / f"{slug}.md").read_text(encoding="utf-8").splitlines()
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)

    d.add_paragraph(title, style="Title")
    p = d.add_paragraph("A how-to guide from Grant Builds With AI · ")
    add_link(p, "grantbuildswithai.com", SITE)
    p.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    para = []

    def flush():
        if para:
            add_inline(d.add_paragraph(), " ".join(para))
            para.clear()

    # Lines that only make sense on the website.
    web_only = {
        " Your checkmarks are saved in this browser.": "",
        "as the website you're reading": "as my website, grantbuildswithai.com",
    }
    for line in md:
        s = line.strip()
        for old, new in web_only.items():
            s = s.replace(old, new)
        if s.startswith("Download the Word version"):
            continue  # this file *is* the Word version
        if not s:
            flush()
        elif s.startswith("### "):
            flush(); d.add_paragraph(s[4:], style="Heading 3")
        elif s.startswith("## "):
            flush(); d.add_paragraph(s[3:], style="Heading 2")
        elif s.startswith("- [ ] "):
            flush(); add_inline(d.add_paragraph("☐ ", style="List Bullet"), s[6:])
        elif s.startswith("- "):
            flush(); add_inline(d.add_paragraph(style="List Bullet"), s[2:])
        elif s.startswith(">"):
            flush()
            if s.strip("> "):
                add_inline(d.add_paragraph(style="Quote"), s.lstrip("> "))
        else:
            para.append(s)
    flush()
    d.save(HERE / out)
    print(f"wrote {out}")


if __name__ == "__main__":
    convert(*sys.argv[1:4])
